import csv, io, re, zipfile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from .search import normalize

COLUMNS={'nom':'name','name':'name','code':'code','plu':'code','synonymes':'keywords','mots cles':'keywords','synonymes mots cles':'keywords','keywords':'keywords','categorie':'category','category':'category','image':'image','photo':'image','precision':'note','note':'note'}

def parse_file(file):
    try: return _parse_file(file)
    except (zipfile.BadZipFile,InvalidFileException,ParseError,KeyError,TypeError,csv.Error,UnicodeError) as exc:
        raise ValueError('Fichier illisible. Enregistre-le de nouveau en CSV ou Excel (.xlsx).') from exc

def _parse_file(file):
    raw=file.read(5*1024*1024+1)
    if len(raw)>5*1024*1024: raise ValueError('Le fichier doit faire moins de 5 Mo.')
    filename=file.filename.lower()
    if filename.endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist())>30*1024*1024: raise ValueError('Classeur trop volumineux après décompression.')
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        sheet=book.active
        if sheet.max_row and sheet.max_row>10001: raise ValueError('Maximum 10 000 produits par import.')
        rows=[]
        for row in sheet.iter_rows():
            values=[]
            for cell in row:
                if cell.data_type=='f': raise ValueError('Remplace les formules par leurs valeurs avant l’import.')
                v=cell.value
                if isinstance(v,(int,float)):
                    if v!=int(v): v=str(v)
                    elif re.fullmatch('0+',cell.number_format or ''): v=str(int(v)).zfill(len(cell.number_format))
                    else: v=str(int(v))
                values.append('' if v is None else str(v).strip())
            rows.append(values)
            if len(rows)>10001: raise ValueError('Maximum 10 000 produits par import.')
        book.close()
    elif filename.endswith('.csv'):
        try: text=raw.decode('utf-8-sig')
        except UnicodeDecodeError: text=raw.decode('cp1252')
        try: dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
        except csv.Error: dialect=csv.excel
        rows=list(csv.reader(io.StringIO(text),dialect))
    else: raise ValueError('Utilise un fichier .csv ou .xlsx.')
    if not rows: raise ValueError('Le fichier est vide.')
    headers=[COLUMNS.get(normalize(h)) for h in rows[0]]
    if 'name' not in headers or 'code' not in headers: raise ValueError('Les colonnes nom et code sont obligatoires.')
    if len([h for h in headers if h])!=len(set(h for h in headers if h)): raise ValueError('Une colonne est présente plusieurs fois.')
    result=[]; seen=set(); errors=[]
    for i,row in enumerate(rows[1:],2):
        if not any(str(v).strip() for v in row): continue
        item={h:str(row[j]).strip() for j,h in enumerate(headers) if h and j<len(row)}
        code=item.get('code','')
        if not code or not item.get('name'): errors.append(f'Ligne {i} : nom ou code manquant.'); continue
        if code in seen: errors.append(f'Ligne {i} : code {code} en double.'); continue
        seen.add(code)
        item.update(demo=False,active=True)
        result.append(item)
    if len(result)>10000: errors.append('Maximum 10 000 produits par import.')
    if errors: raise ValueError('\n'.join(errors[:30]))
    if not result: raise ValueError('Aucun produit à importer.')
    return result
